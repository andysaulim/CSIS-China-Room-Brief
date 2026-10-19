"""Google Sheets access for Brian's sheets.

With a service account (GOOGLE_SERVICE_ACCOUNT_JSON, the key file's JSON) the
run can read sheets that are shared only with that account's email, and write
the Talk Around Town approvals tab. Without one, reading falls back to the
public xlsx export, which works for sheets shared as "Anyone with the link".

    xlsx(url, dest)              download any sheet as .xlsx (tracker.load reads it)
    rows(url, tab=None)          header-keyed rows from one tab
    append(url, tab, headers, rows)   add rows to a tab, creating it with headers if needed
"""

from __future__ import annotations

import json
import os
import re
from urllib.parse import quote

import requests

SCOPES = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive.readonly"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_token = None


def sheet_id(url: str) -> str:
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]{20,})", url or "")
    return m.group(1) if m else ""


def _auth() -> dict | None:
    """Bearer header from the service account, or None when there is no key."""
    global _token
    raw = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()
    if not raw:
        return None
    if _token is None:
        from google.auth.transport.requests import Request
        from google.oauth2 import service_account
        creds = service_account.Credentials.from_service_account_info(json.loads(raw), scopes=SCOPES)
        creds.refresh(Request())
        _token = creds.token
    return {"Authorization": f"Bearer {_token}"}


def account_email() -> str:
    raw = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()
    try:
        return json.loads(raw).get("client_email", "") if raw else ""
    except ValueError:
        return ""


def xlsx(url: str, dest: str) -> str | None:
    """Save the sheet at url as an .xlsx file. Prints why when it cannot."""
    sid = sheet_id(url)
    auth = _auth() if sid else None
    try:
        if sid and auth:
            r = requests.get(f"https://www.googleapis.com/drive/v3/files/{sid}/export",
                             params={"mimeType": XLSX}, headers=auth, timeout=60)
            if r.status_code in (403, 404):
                print(f"sheet {sid[:6]}…: the service account cannot open it; share it with {account_email()}")
        else:
            export = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx" if sid else url
            r = requests.get(export, timeout=60)
    except requests.RequestException as e:
        print(f"sheet download: {type(e).__name__}")
        return None
    if not r.ok or r.content[:2] != b"PK":
        print(f"sheet download failed: HTTP {r.status_code}"
              + ("" if auth else " (no service account; is the sheet shared as Anyone with the link?)"))
        return None
    with open(dest, "wb") as f:
        f.write(r.content)
    return dest


def _api(sid: str, path: str = "", method: str = "GET", **kw):
    auth = _auth()
    if not auth:
        raise RuntimeError("GOOGLE_SERVICE_ACCOUNT_JSON is not set")
    r = requests.request(method, f"https://sheets.googleapis.com/v4/spreadsheets/{sid}{path}",
                         headers=auth, timeout=60, **kw)
    if not r.ok:
        raise RuntimeError(f"Sheets API {r.status_code}: {r.text[:200]}")
    return r.json()


def tabs(url: str) -> list[str]:
    data = _api(sheet_id(url), params={"fields": "sheets.properties.title"})
    return [s["properties"]["title"] for s in data.get("sheets", [])]


def rows(url: str, tab: str | None = None) -> list[dict]:
    """Rows of one tab as dicts keyed by the header row. Uses the API when it can, the xlsx otherwise."""
    if _auth():
        name = tab or tabs(url)[0]
        values = _api(sheet_id(url), f"/values/{quote(name)}").get("values", [])
    else:
        import tempfile

        import openpyxl
        path = xlsx(url, os.path.join(tempfile.gettempdir(), f"sheet_{sheet_id(url)[:8]}.xlsx"))
        if not path:
            return []
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[tab] if tab and tab in wb.sheetnames else wb[wb.sheetnames[0]]
        values = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    values = [v for v in values if any(str(c).strip() for c in v)]
    if not values:
        return []
    head = [str(h).strip() for h in values[0]]
    return [{head[i]: (r[i] if i < len(r) else "") for i in range(len(head)) if head[i]} for r in values[1:]]


def append(url: str, tab: str, headers: list[str], new_rows: list[list]) -> None:
    """Append rows to tab, creating the tab (with headers) when it does not exist."""
    sid = sheet_id(url)
    if tab not in tabs(url):
        _api(sid, ":batchUpdate", "POST", json={"requests": [{"addSheet": {"properties": {"title": tab}}}]})
        _api(sid, f"/values/{quote(tab)}!A1:append", "POST", params={"valueInputOption": "RAW"},
             json={"values": [headers]})
    if new_rows:
        _api(sid, f"/values/{quote(tab)}!A1:append", "POST",
             params={"valueInputOption": "RAW", "insertDataOption": "INSERT_ROWS"}, json={"values": new_rows})
